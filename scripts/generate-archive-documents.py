"""Regenerate the fictional Archive DOCX and PDF sample documents."""

from pathlib import Path

from docx import Document
from docx.shared import Inches
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


ASSETS = Path(__file__).resolve().parents[1] / "src" / "assets"


def make_docx():
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    document.add_heading("Pawford research visit", 0)
    document.add_paragraph("Field notes · 2 June 2026", style="Subtitle")
    document.add_heading("Purpose", level=1)
    document.add_paragraph(
        "Compare the Whiskerfield household register with the surviving name index "
        "and note where a second source is still needed."
    )
    document.add_heading("Observations", level=1)
    for item in (
        "The index points to page 12 for Barnaby, Daisy, Silver, and Luna Whiskerfield.",
        "The register image is legible, but relationships should be checked against original records.",
        "Daisy's letter may help identify photographs associated with this household.",
    ):
        document.add_paragraph(item, style="List Bullet")
    document.add_heading("Follow-up", level=1)
    table = document.add_table(rows=1, cols=2)
    table.style = "Light Shading Accent 1"
    table.rows[0].cells[0].text = "Item"
    table.rows[0].cells[1].text = "Next step"
    for name, action in (
        ("Household register", "Compare page 12 against the index transcription"),
        ("Birth announcement", "Confirm Silver's date and place"),
        ("Family letter", "Identify the photographs Daisy describes"),
    ):
        cells = table.add_row().cells
        cells[0].text = name
        cells[1].text = action
    document.add_paragraph("Sample research material for the GeneoGraph prototype.")
    document.save(ASSETS / "pawford_visit_notes.docx")


def make_pdf():
    target = ASSETS / "silver_luna_marriage_packet.pdf"
    styles = getSampleStyleSheet()
    document = SimpleDocTemplate(str(target), pagesize=A4, title="Silver and Luna marriage packet")
    story = [
        Paragraph("Silver &amp; Luna: marriage research packet", styles["Title"]),
        Spacer(1, 14),
        Paragraph("Meowbridge · 18 August 2023", styles["Heading2"]),
        Paragraph(
            "This packet summarizes the sample marriage extract and records the questions "
            "to resolve before treating the date and place as verified.",
            styles["BodyText"],
        ),
        Spacer(1, 16),
        Paragraph("Evidence checklist", styles["Heading2"]),
    ]
    table = Table(
        [
            ["Record", "Observation", "Status"],
            ["Marriage extract", "Names Silver and Luna; Meowbridge", "Sample image"],
            ["Household index", "Both names appear on page 12", "Cross-reference"],
            ["Original register", "Request a full citation and image", "Open"],
        ],
        colWidths=[120, 270, 90],
        repeatRows=1,
    )
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#d8e9e1")),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 8),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
    ]))
    story.extend([
        table,
        Spacer(1, 16),
        Paragraph("Research note", styles["Heading2"]),
        Paragraph(
            "The extract is illustrative. Preserve uncertainty until an original source "
            "with full provenance is available.",
            styles["BodyText"],
        ),
        PageBreak(),
        Paragraph("Research questions and next steps", styles["Title"]),
        Spacer(1, 18),
        Paragraph("1. Where is the original marriage entry held?", styles["Heading2"]),
        Paragraph("Record repository, collection, volume, page, and stable reference when found.", styles["BodyText"]),
        Spacer(1, 12),
        Paragraph("2. Do both names match other family documents?", styles["Heading2"]),
        Paragraph("Compare spelling and dates with the Pawford household index and personal records.", styles["BodyText"]),
        Spacer(1, 12),
        Paragraph("3. Which place should be cited?", styles["Heading2"]),
        Paragraph("Use Meowbridge as the working place until the original entry confirms it.", styles["BodyText"]),
    ])
    document.build(story)


if __name__ == "__main__":
    make_docx()
    make_pdf()
